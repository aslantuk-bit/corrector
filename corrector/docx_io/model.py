"""Документ Word как список абзацев с картой фрагментов форматирования (run)."""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.document import Document as _Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

BROKEN_MESSAGE = (
    "Файл повреждён, защищён паролем или это не документ Word (.docx). "
    "Старый формат .doc откройте в Word и сохраните как .docx."
)


class DocxError(Exception):
    pass


@dataclass
class RunSpan:
    run: Run
    start: int
    end: int


@dataclass
class Para:
    index: int
    text: str
    runs: list[RunSpan]
    where: str
    paragraph: Paragraph


@dataclass
class ParagraphBlock:
    index: int


@dataclass
class TableBlock:
    rows: list[list[list]]  # строки → ячейки → блоки ячейки


Block = ParagraphBlock | TableBlock


@dataclass
class DocumentModel:
    path: Path
    document: _Document
    paragraphs: list[Para]
    blocks: list = field(default_factory=list)  # порядок и структура тела: абзацы и таблицы (с вложенными)

    def texts(self) -> list[str]:
        return [p.text for p in self.paragraphs]


def load(path: Path) -> DocumentModel:
    path = Path(path)
    if not path.is_file():
        raise DocxError(f"Файл не найден: {path}")
    try:
        document = docx.Document(str(path))
    except (PackageNotFoundError, KeyError, ValueError, zipfile.BadZipFile, OSError) as error:
        raise DocxError(BROKEN_MESSAGE) from error
    paragraphs: list[Para] = []

    def add(paragraph: Paragraph, where: str, blocks: list) -> None:
        spans, parts, position = [], [], 0
        for item in paragraph.iter_inner_content():
            if isinstance(item, Run):
                runs = [item]
            elif isinstance(item, Hyperlink):
                runs = list(item.runs)
            else:
                runs = []
            for run in runs:
                text = run.text
                if text:
                    spans.append(RunSpan(run, position, position + len(text)))
                    parts.append(text)
                    position += len(text)
        paragraphs.append(Para(len(paragraphs), "".join(parts), spans, where, paragraph))
        blocks.append(ParagraphBlock(len(paragraphs) - 1))
        # надписи (текстовые поля) внутри абзаца: их абзацы идут следом
        for content in paragraph._p.iter(qn("w:txbxContent")):
            for p_element in content.findall(qn("w:p")):
                add(Paragraph(p_element, paragraph._parent), "textbox", blocks)

    def walk(container, where: str, blocks: list) -> None:
        for item in container.iter_inner_content():
            if isinstance(item, Paragraph):
                add(item, where, blocks)
            elif isinstance(item, Table):
                rows: list[list[list]] = []
                seen = []
                for row in item.rows:
                    cells: list[list] = []
                    for cell in row.cells:
                        if any(cell._tc is tc for tc in seen):
                            continue
                        seen.append(cell._tc)
                        cell_blocks: list = []
                        walk(cell, "table", cell_blocks)
                        cells.append(cell_blocks)
                    rows.append(cells)
                blocks.append(TableBlock(rows))

    body_blocks: list = []
    walk(document, "body", body_blocks)
    for section in document.sections:
        # колонтитул без собственного определения читать нельзя: python-docx создал бы его и изменил документ
        if not section.header.is_linked_to_previous:
            for paragraph in section.header.paragraphs:
                if paragraph.text.strip():
                    add(paragraph, "header", body_blocks)
        if not section.footer.is_linked_to_previous:
            for paragraph in section.footer.paragraphs:
                if paragraph.text.strip():
                    add(paragraph, "footer", body_blocks)
    return DocumentModel(path, document, paragraphs, body_blocks)
