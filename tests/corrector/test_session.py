import hashlib
import json

import docx
import pytest

from corrector.app.session import Session
from corrector.core.issue import Category, Level
from corrector.core.userdict import UserDictionary
from corrector.engines import factory
from tests.corrector.conftest import make_docx


@pytest.fixture(scope="module")
def engines(tmp_path_factory):
    base = tmp_path_factory.mktemp("движки")
    engines = factory.build_engines(UserDictionary(base / "словарь.txt"), no_lt=True, user_dir=base)
    yield engines
    engines.close()


def make_session(tmp_path, engines, body):
    path = make_docx(tmp_path / "акт.docx", body=body)
    session = Session(engines, tmp_path)
    session.open(path)
    session.check()
    return session


def test_open_check_lists_issues_in_order(tmp_path, engines):
    session = make_session(tmp_path, engines, ["Соттын шешімі заңды.", "Истец подал ходатайтсво в суд ,  решил."])
    fragments = [(i.issue.paragraph, session.fragment(i)) for i in session.open_items()]
    assert fragments[0] == (0, "Соттын")
    assert (1, "ходатайтсво") in fragments and (1, "  ") in fragments
    assert any(i.issue.rule_id == "typo.space_before_punct" for i in session.open_items())
    assert session.counts()[Level.ERROR] >= 2


def test_apply_shifts_later_issues_and_closes_overlapping(tmp_path, engines):
    session = make_session(tmp_path, engines, ["Истец подал ходатайтсво в суд ,  решил."])
    spelling = next(i for i in session.open_items() if session.fragment(i) == "ходатайтсво")
    comma = next(i for i in session.open_items() if i.issue.rule_id == "typo.space_before_punct")
    session.apply(spelling.id, "заявление")
    assert spelling.status == "applied"
    assert session.model.paragraphs[0].text == "Истец подал заявление в суд ,  решил."
    assert session.fragment(comma) == " " and session.model.paragraphs[0].text[comma.issue.end] == ","
    session.apply(comma.id, "")
    assert session.model.paragraphs[0].text == "Истец подал заявление в суд,  решил."


def test_skip_dictionary_and_disable_rule(tmp_path, engines):
    session = make_session(tmp_path, engines, ["Соттын шешімі. Соттын қаулысы.", "Истец подал ходатайтсво в суд , решил."])
    first, second = [i for i in session.open_items() if session.fragment(i) == "Соттын"]
    session.skip(first.id)
    assert first.status == "skipped" and second.status == "open"
    session.add_to_dictionary(second.id)
    assert second.status == "closed"
    assert "Соттын" in (tmp_path / "словарь.txt").read_text(encoding="utf-8")
    comma = next(i for i in session.open_items() if i.issue.rule_id == "typo.space_before_punct")
    session.disable_rule(comma.id)
    assert comma.status == "closed"
    assert comma.issue.rule_id in json.loads((tmp_path / "настройки.json").read_text(encoding="utf-8"))["disabled_rules"]
    session.check()
    assert all(i.issue.rule_id != "typo.space_before_punct" for i in session.open_items())
    assert all(session.fragment(i) != "Соттын" for i in session.open_items())


def test_category_filter(tmp_path, engines):
    session = make_session(tmp_path, engines, ["Истец подал ходатайтсво в суд ,  решил."])
    assert any(i.issue.category is Category.TYPOGRAPHY for i in session.visible_items())
    session.set_category("typography", False)
    assert not any(i.issue.category is Category.TYPOGRAPHY for i in session.visible_items())
    assert any(i.issue.category is Category.SPELLING for i in session.visible_items())


def test_save_with_comments_keeps_original(tmp_path, engines):
    session = make_session(tmp_path, engines, ["Истец подал ходатайтсво в суд ,  решил."])
    original = tmp_path / "акт.docx"
    before = hashlib.sha256(original.read_bytes()).hexdigest()
    spelling = next(i for i in session.open_items() if session.fragment(i) == "ходатайтсво")
    session.apply(spelling.id, "ходатайство")
    target = session.save(with_comments=True)
    assert target == tmp_path / "акт_проверено.docx" and target.exists()
    assert hashlib.sha256(original.read_bytes()).hexdigest() == before
    saved = docx.Document(str(target))
    assert saved.paragraphs[0].text == "Истец подал ходатайство в суд ,  решил."
    comments = list(saved.comments)
    assert len(comments) == len(session.visible_items()) >= 2
    assert comments[0].author == "Корректор" and "Типографика" in comments[0].text
