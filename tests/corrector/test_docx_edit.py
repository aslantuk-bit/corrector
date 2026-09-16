import hashlib

import docx

from corrector.docx_io import edit, model
from tests.corrector.conftest import make_docx


def reload(path):
    return model.load(path)


def test_replacement_inside_one_run_keeps_formatting(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=[[("Истец подал ", {}), ("ходатайтсво", {"bold": True}), (" в суд.", {})]])
    doc = reload(path)
    para = doc.paragraphs[0]
    start = para.text.index("ходатайтсво")
    delta = edit.apply_replacement(para, start, start + len("ходатайтсво"), "ходатайство")
    assert delta == 0
    assert para.text == "Истец подал ходатайство в суд."
    assert [(s.start, s.end, s.run.bold) for s in para.runs] == [(0, 12, None), (12, 23, True), (23, 30, None)]
    target = tmp_path / "б.docx"
    edit.save_copy(doc, target)
    again = reload(target)
    assert again.paragraphs[0].text == "Истец подал ходатайство в суд."
    assert again.paragraphs[0].runs[1].run.bold is True
    delta = edit.apply_replacement(para, 12, 23, "заявление")
    assert delta == -2 and para.text == "Истец подал заявление в суд."


def test_replacement_across_run_boundary(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=[[("в соответ", {"italic": True}), ("ствие с законом", {})]])
    doc = reload(path)
    para = doc.paragraphs[0]
    edit.apply_replacement(para, 0, len("в соответствие с"), "в соответствии с")
    assert para.text == "в соответствии с законом"
    assert "".join(s.run.text for s in para.runs) == para.text
    assert para.runs[0].run.italic is True


def test_replacement_keeps_tabs(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=[[("а\tб ошибка", {})]])
    doc = reload(path)
    para = doc.paragraphs[0]
    start = para.text.index("ошибка")
    edit.apply_replacement(para, start, start + 6, "верно")
    assert para.text == "а\tб верно"
    edit.save_copy(doc, tmp_path / "б.docx")
    assert reload(tmp_path / "б.docx").paragraphs[0].text == "а\tб верно"


def test_split_run_and_comment_on_fragment(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=[[("Истец подал ходатайтсво в суд.", {"bold": True})]])
    doc = reload(path)
    para = doc.paragraphs[0]
    start = para.text.index("ходатайтсво")
    assert edit.add_comment(doc, para, start, start + 11, "Орфография: возможно, «ходатайство»") is True
    assert para.text == "Истец подал ходатайтсво в суд."
    assert [(s.start, s.end) for s in para.runs] == [(0, 12), (12, 23), (23, 30)]
    assert all(s.run.bold is True for s in para.runs)
    edit.save_copy(doc, tmp_path / "б.docx")
    saved = docx.Document(str(tmp_path / "б.docx"))
    comments = list(saved.comments)
    assert len(comments) == 1 and comments[0].author == "Корректор" and "ходатайство" in comments[0].text
    assert saved.paragraphs[0].text == "Истец подал ходатайтсво в суд."


def test_comment_in_header_is_skipped(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["Тело"], header="Шапка с ошибкой")
    doc = reload(path)
    header = [p for p in doc.paragraphs if p.where == "header"][0]
    assert edit.add_comment(doc, header, 0, 5, "м") is False


def test_output_path_and_original_untouched(tmp_path):
    path = make_docx(tmp_path / "акт.docx", body=["Текст"])
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    first = edit.output_path(path)
    assert first == tmp_path / "акт_проверено.docx"
    first.write_bytes("занято".encode("utf-8"))
    assert edit.output_path(path) == tmp_path / "акт_проверено (2).docx"
    doc = reload(path)
    edit.apply_replacement(doc.paragraphs[0], 0, 5, "Новый")
    saved = edit.save_copy(doc, edit.output_path(path))
    assert saved.name == "акт_проверено (2).docx" and saved.exists()
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_comment_in_text_box_does_not_crash(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["Основной текст"])
    from docx.oxml import parse_xml
    document = docx.Document(str(path))
    run = document.paragraphs[0].add_run()
    xml = ('<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
           '<v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Соттын шешімі</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')
    run._r.append(parse_xml(xml))
    document.save(str(path))
    doc = reload(path)
    box = [p for p in doc.paragraphs if p.where == "textbox"][0]
    result = edit.add_comment(doc, box, 0, 6, "м")
    assert result in (True, False)
    edit.save_copy(doc, tmp_path / "б.docx")
    assert docx.Document(str(tmp_path / "б.docx")).paragraphs[0].text == "Основной текст"
