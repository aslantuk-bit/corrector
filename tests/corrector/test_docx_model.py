import pytest

from corrector.docx_io import model
from tests.corrector.conftest import make_docx


def test_load_body_paragraphs_with_run_map(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["Первый абзац.", [("Жирный", {"bold": True}), (" и обычный", {})]])
    doc = model.load(path)
    assert doc.texts() == ["Первый абзац.", "Жирный и обычный"]
    second = doc.paragraphs[1]
    assert [(s.start, s.end, s.run.text) for s in second.runs] == [(0, 6, "Жирный"), (6, 16, " и обычный")]
    assert second.runs[0].run.bold is True
    assert second.where == "body"


def test_hyperlink_runs_are_part_of_text(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=[[("См. ", {}), ("hyperlink", "сайт суда", "https://sud.kz"), (" здесь", {})]])
    doc = model.load(path)
    assert doc.texts() == ["См. сайт суда здесь"]
    assert [s.run.text for s in doc.paragraphs[0].runs] == ["См. ", "сайт суда", " здесь"]


def test_tabs_and_breaks_keep_offsets(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=[[("а\tб\nв", {}), ("г", {})]])
    doc = model.load(path)
    assert doc.texts() == ["а\tб\nв" + "г"]
    assert [(s.start, s.end) for s in doc.paragraphs[0].runs] == [(0, 5), (5, 6)]


def test_tables_headers_footers_in_order(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["Тело"], table=[["Я1", "Я2"], ["Я3", "Я4"]], header="Шапка", footer="Подвал")
    doc = model.load(path)
    assert [(p.text, p.where) for p in doc.paragraphs] == [
        ("Тело", "body"), ("Я1", "table"), ("Я2", "table"), ("Я3", "table"), ("Я4", "table"),
        ("Шапка", "header"), ("Подвал", "footer"),
    ]
    assert [p.index for p in doc.paragraphs] == list(range(7))


def test_merged_cells_are_read_once(tmp_path):
    path = make_docx(tmp_path / "а.docx", table=[["Один", "Два"]])
    import docx
    document = docx.Document(str(path))
    table = document.tables[0]
    table.cell(0, 0).merge(table.cell(0, 1))
    document.save(str(path))
    doc = model.load(path)
    assert len([p for p in doc.paragraphs if p.where == "table"]) == 2


def test_broken_file_raises_docx_error(tmp_path):
    bad = tmp_path / "б.docx"
    bad.write_bytes("это не документ".encode("utf-8"))
    with pytest.raises(model.DocxError) as info:
        model.load(bad)
    assert "не документ Word" in str(info.value)


def test_missing_file_raises_docx_error(tmp_path):
    with pytest.raises(model.DocxError):
        model.load(tmp_path / "нет.docx")


def test_load_does_not_create_headers(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["Тело"])
    doc = model.load(path)
    assert doc.texts() == ["Тело"]
    assert doc.document.sections[0].header.is_linked_to_previous is True
    assert doc.document.sections[0].footer.is_linked_to_previous is True


def test_layout_blocks_keep_table_structure(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["До"], table=[["Я1", "Я2"], ["Я3", "Я4"]])
    doc = model.load(path)
    kinds = [type(b).__name__ for b in doc.blocks]
    assert kinds == ["ParagraphBlock", "TableBlock"]
    table = doc.blocks[1]
    assert len(table.rows) == 2 and len(table.rows[0]) == 2
    first_cell = table.rows[0][0]
    assert [type(b).__name__ for b in first_cell] == ["ParagraphBlock"]
    assert doc.paragraphs[first_cell[0].index].text == "Я1"


def test_nested_table_and_paragraph_order(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["До"], table=[["Ячейка"]])
    import docx as _docx
    document = _docx.Document(str(path))
    cell = document.tables[0].cell(0, 0)
    inner = cell.add_table(rows=1, cols=1)
    inner.cell(0, 0).paragraphs[0].add_run("Вложенная")
    document.add_paragraph("После")
    document.save(str(path))
    doc = model.load(path)
    assert [t for t in doc.texts() if t] == ["До", "Ячейка", "Вложенная", "После"]  # python-docx добавляет пустой абзац после таблицы
    cell_blocks = doc.blocks[1].rows[0][0]
    assert [type(b).__name__ for b in cell_blocks][:2] == ["ParagraphBlock", "TableBlock"]


def test_text_boxes_are_read(tmp_path):
    path = make_docx(tmp_path / "а.docx", body=["Основной текст"])
    import docx as _docx
    document = _docx.Document(str(path))
    run = document.paragraphs[0].add_run()
    from docx.oxml import parse_xml
    xml = ('<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml">'
           '<v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Текст в надписи</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')
    run._r.append(parse_xml(xml))
    document.save(str(path))
    doc = model.load(path)
    assert [(p.text, p.where) for p in doc.paragraphs] == [("Основной текст", "body"), ("Текст в надписи", "textbox")]
