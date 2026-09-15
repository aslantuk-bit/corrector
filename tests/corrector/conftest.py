from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _fill(paragraph, content) -> None:
    if isinstance(content, str):
        paragraph.add_run(content)
        return
    for piece in content:
        if piece[0] == "hyperlink":
            _add_hyperlink(paragraph, piece[1], piece[2])
        else:
            text, props = piece
            run = paragraph.add_run(text)
            for key, value in props.items():
                setattr(run, key, value)


def make_docx(path: Path, body=(), table=None, header=None, footer=None) -> Path:
    document = docx.Document()
    for content in body:
        _fill(document.add_paragraph(), content)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell_text in enumerate(row):
                _fill(grid.cell(r, c).paragraphs[0], cell_text)
    if header is not None:
        _fill(document.sections[0].header.paragraphs[0], header)
    if footer is not None:
        _fill(document.sections[0].footer.paragraphs[0], footer)
    document.save(str(path))
    return path
